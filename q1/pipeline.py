from __future__ import annotations

import argparse
import json
import logging
import os
import re
import shutil
import subprocess
import sys
import time
import traceback
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import numpy as np

from .alignment import align_text_words, make_time_grid, overlap_pool, standardize_aligned
from .config import Q1Config, dump_yaml, load_config
from .extractors import ASRAligner, AudioExtractor, TextExtractor, VisionExtractor
from .utils import (
    atomic_pickle_dump,
    command_version,
    configure_logging,
    environment_snapshot,
    extract_audio,
    probe_media,
    set_deterministic,
    sha256_file,
    verify_model_manifest,
    write_csv,
    write_json,
)


def _normal_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def _string_id(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _column_index(headers: list[Any], names: list[str]) -> int:
    normalized = {_normal_header(value): index for index, value in enumerate(headers)}
    for name in names:
        key = _normal_header(name)
        if key in normalized:
            return normalized[key]
    raise KeyError(f"Could not find any of {names} in columns {headers}")


def read_label_table(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = list(worksheet.iter_rows(values_only=True))
    header_row = next((index for index, row in enumerate(rows) if any(value is not None for value in row)), None)
    if header_row is None:
        raise ValueError(f"Empty label table: {path}")
    headers = list(rows[header_row])
    indices = {
        "video_id": _column_index(headers, ["video_id", "video id"]),
        "clip_id": _column_index(headers, ["clip_id", "clip id"]),
        "text": _column_index(headers, ["text", "raw_text", "raw text"]),
        "label": _column_index(headers, ["label", "regression_label", "regression label"]),
        "annotation": _column_index(headers, ["annotation", "classification", "classification_label"]),
    }
    records = []
    for row in rows[header_row + 1 :]:
        if not any(value is not None for value in row):
            continue
        video_id = _string_id(row[indices["video_id"]])
        clip_id = _string_id(row[indices["clip_id"]])
        if not video_id or not clip_id:
            continue
        label_value = row[indices["label"]]
        try:
            label = float(label_value) if label_value is not None else None
        except (TypeError, ValueError):
            label = None
        records.append(
            {
                "video_id": video_id,
                "clip_id": clip_id,
                "id": f"{video_id}$_${clip_id}",
                "raw_text": str(row[indices["text"]] or ""),
                "label": label,
                "annotation": str(row[indices["annotation"]] or ""),
            }
        )
    if len(records) != 100:
        raise ValueError(f"Expected 100 label rows, found {len(records)} in {path}")
    ids = [record["id"] for record in records]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate sample IDs in label table")
    return records


def build_manifest(config: Q1Config, logger: logging.Logger) -> list[dict[str, Any]]:
    label_file = config.path_for("label_file")
    video_root = config.path_for("video_root")
    labels = read_label_table(label_file)
    video_map = {
        (path.parent.name, path.stem): path
        for path in video_root.rglob("*.mp4")
    }
    if len(video_map) != 100:
        raise ValueError(f"Expected exactly 100 MP4 files under {video_root}, found {len(video_map)}")
    records = []
    for label in labels:
        key = (label["video_id"], label["clip_id"])
        video_path = video_map.get(key)
        if video_path is None:
            raise FileNotFoundError(f"No video for {label['id']}: {key}")
        media = probe_media(video_path)
        if media["duration"] <= 0:
            raise ValueError(f"Invalid video duration for {video_path}")
        source_duration_range = config.get("source_duration_range_s", [0.0, float("inf")])
        container_duration = float(media["duration_container"])
        if not (float(source_duration_range[0]) - 1e-4 <= container_duration <= float(source_duration_range[1]) + 1e-4):
            raise ValueError(
                f"Container duration {container_duration:.6f}s outside configured source range "
                f"{source_duration_range} for {video_path}"
            )
        record = {
            **label,
            "video_path": str(video_path),
            "video_path_relative": str(video_path.relative_to(config.project_root)),
            "duration": float(media["duration_alignment"]),  # internal compatibility alias
            "duration_container": float(media["duration_container"]),
            "duration_format": float(media["duration_format"]),
            "duration_mvhd": float(media["duration_mvhd"]),
            "duration_video_stream": float(media["duration_video_stream"]),
            "duration_audio_stream": float(media["duration_audio_stream"]),
            "duration_video_edit": float(media["duration_video_edit"]),
            "duration_audio_edit": float(media["duration_audio_edit"]),
            "duration_decoded_video": float(media["duration_decoded_video"]),
            "duration_decoded_audio": float(media["duration_decoded_audio"]),
            "duration_decoded": float(media["duration_decoded"]),
            "duration_alignment": float(media["duration_alignment"]),
            "duration_source": str(media["duration_source"]),
            "duration_start_time": float(media["duration_start_time"]),
            "video_edit_list_entries": int(media["video_edit_list_entries"]),
            "audio_edit_list_entries": int(media["audio_edit_list_entries"]),
            "video_edit_list_present": int(media["video_edit_list_present"]),
            "audio_edit_list_present": int(media["audio_edit_list_present"]),
            "video_fps": float(media["video_fps"]),
            "video_width": int(media["video_width"]),
            "video_height": int(media["video_height"]),
            "video_codec": media["video_codec"],
            "audio_codec": media["audio_codec"],
            "audio_sample_rate_original": media["audio_sample_rate_original"],
            "has_audio": int(media["has_audio"]),
            "sha256": sha256_file(video_path),
            "feature_status": "indexed",
        }
        record.update(_duration_endpoint_audit(media, config))
        records.append(record)
    logger.info("Manifest validated: %d labels, %d videos, %d unique IDs", len(labels), len(video_map), len({x['id'] for x in records}))
    return records


def _safe_name(sample_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", sample_id)


def _duration_endpoint_audit(media: dict[str, Any], config: Q1Config) -> dict[str, Any]:
    """Describe and validate the edit-list tail between alignment and decoding."""
    alignment = float(media.get("duration_alignment", 0.0))
    decoded_video = float(media.get("duration_decoded_video_end_pts", media.get("duration_decoded_video", 0.0)))
    decoded_audio = float(media.get("duration_decoded_audio_end_pts", media.get("duration_decoded_audio", 0.0)))
    video_gap = alignment - decoded_video if decoded_video > 0 else 0.0
    audio_gap = alignment - decoded_audio if decoded_audio > 0 else 0.0
    decoded_max = max(decoded_video, decoded_audio)
    max_gap = alignment - decoded_max if decoded_max > 0 else 0.0
    output_fps = max(float(config.get("video_output_fps", 25.0)), 1.0)
    frame_count = max(int(config.get("duration_endpoint_video_frame_count", 2)), 1)
    video_tolerance = max(
        float(config.get("duration_endpoint_video_min_tolerance_s", 0.10)),
        frame_count / output_fps,
    )
    audio_tolerance = float(config.get("duration_endpoint_audio_tolerance_s", 0.25))
    negative_tolerance = float(config.get("duration_endpoint_negative_tolerance_s", 0.02))
    issues: list[str] = []
    if decoded_video > 0 and video_gap < -negative_tolerance:
        issues.append("decoded_video_end_exceeds_alignment")
    if decoded_audio > 0 and audio_gap < -negative_tolerance:
        issues.append("decoded_audio_end_exceeds_alignment")
    if decoded_video > 0 and video_gap > video_tolerance:
        issues.append("video_edit_list_tail_exceeds_tolerance")
    if decoded_audio > 0 and audio_gap > audio_tolerance:
        issues.append("audio_edit_list_tail_exceeds_tolerance")
    if issues:
        status = "manual_review"
    elif max(video_gap, audio_gap) > 0:
        status = "accepted_edit_list_tail"
    else:
        status = "accepted_decoded_endpoint"
    rule = (
        "alignment=video_edit_list_presentation; positive decoded tail is outside decoded spans and remains mask=0; "
        f"video_gap<=max({frame_count}/video_output_fps,{video_tolerance:.2f}s); "
        f"audio_gap<={audio_tolerance:.2f}s; decoded_end may not exceed alignment by>{negative_tolerance:.2f}s"
    )
    explanation = (
        "The alignment axis starts at the edit-list presentation zero. Decoded frame/packet endpoints can finish "
        "earlier because of frame-center/packet-boundary quantization or edit-list padding. Positive tails are not "
        "imputed: overlap pooling clips spans to the alignment grid and leaves uncovered bins masked."
    )
    return {
        "duration_decoded_video_start_pts": float(media.get("duration_decoded_video_start_pts", 0.0)),
        "duration_decoded_video_last_pts": float(media.get("duration_decoded_video_last_pts", 0.0)),
        "duration_decoded_video_end_pts": decoded_video,
        "duration_decoded_video_packet_count": int(media.get("duration_decoded_video_packet_count", 0)),
        "duration_decoded_video_packet_duration": float(media.get("duration_decoded_video_packet_duration", 0.0)),
        "duration_decoded_audio_start_pts": float(media.get("duration_decoded_audio_start_pts", 0.0)),
        "duration_decoded_audio_last_pts": float(media.get("duration_decoded_audio_last_pts", 0.0)),
        "duration_decoded_audio_end_pts": decoded_audio,
        "duration_decoded_audio_packet_count": int(media.get("duration_decoded_audio_packet_count", 0)),
        "duration_decoded_audio_packet_duration": float(media.get("duration_decoded_audio_packet_duration", 0.0)),
        "duration_alignment_minus_decoded_video": float(video_gap),
        "duration_alignment_minus_decoded_audio": float(audio_gap),
        "duration_alignment_minus_decoded_max": float(max_gap),
        "duration_endpoint_video_tolerance": float(video_tolerance),
        "duration_endpoint_audio_tolerance": float(audio_tolerance),
        "duration_endpoint_negative_tolerance": float(negative_tolerance),
        "duration_endpoint_status": status,
        "duration_endpoint_issues": ";".join(issues),
        "duration_endpoint_rule": rule,
        "duration_endpoint_explanation": explanation,
    }


def _text_review_conclusion(stats: dict[str, Any]) -> str:
    status = str(stats.get("text_review_status", ""))
    if status == "asr_empty_or_nonlexical_manual_review":
        return "ASR为空或非词汇，词级时间定位不可依赖"
    if status == "asr_text_mismatch_manual_review":
        return "ASR与题目转写失配，保留低置信度回退时间并人工复核"
    if status == "high_fallback_manual_review":
        return "回退比例较高，作为低置信度文本并人工复核"
    if status == "partial_invalid_timestamp_manual_review":
        return "部分词缺少有效时间戳，保留无效原因并人工复核"
    return "可作为常规文本对齐记录"


def _empty_result(config: Q1Config, entry: dict[str, Any], reason: str) -> dict[str, Any]:
    length = config.aligned_length
    text_dim, audio_dim, vision_dim = config.main_feature_dims
    return {
        "text": np.zeros((length, text_dim), dtype=np.float32),
        "audio": np.zeros((length, audio_dim), dtype=np.float32),
        "vision": np.zeros((length, vision_dim), dtype=np.float32),
        "text_mask": np.zeros(length, dtype=np.uint8),
        "text_confidence": np.zeros(length, dtype=np.float32),
        "text_low_confidence_mask": np.zeros(length, dtype=np.uint8),
        "audio_mask": np.zeros(length, dtype=np.uint8),
        "vision_mask": np.zeros(length, dtype=np.uint8),
        "vision_emotion_probs": np.zeros((length, 7), dtype=np.float32),
        "audio_prosody": np.zeros((length, 5), dtype=np.float32),
        "audio_prosody_mask": np.zeros(length, dtype=np.uint8),
        "duration": float(entry["duration_alignment"]),
        "duration_container": float(entry["duration_container"]),
        "duration_alignment": float(entry["duration_alignment"]),
        "time_grid": make_time_grid(float(entry["duration_alignment"]), length),
        "raw_lengths": {"text": 0, "audio": 0, "vision": 0},
        "valid_lengths": {"text": 0, "audio": 0, "vision": 0},
        "stats": {
            "text_valid_ratio": 0.0,
            "text_low_confidence_ratio": 0.0,
            "text_fallback_count": 0,
            "text_fallback_ratio": 0.0,
            "text_lexical_asr_word_count": 0,
            "text_review_status": "processing_failed_manual_review",
            "text_asr_retry_used": 0,
            "text_asr_retry_selected": 0,
            "text_asr_selected_profile": "none",
            "text_asr_attempt_count": 0,
            "audio_valid_ratio": 0.0,
            "vision_valid_ratio": 0.0,
            "vision_face_detection_ratio": 0.0,
            "vision_imputed_ratio": 0.0,
            "failure_reason": reason,
        },
        "raw_audit": {"id": entry["id"], "error": reason},
        "status": "failed",
        "processing_status": "failed",
        "quality_status": "missing",
        "modality_status": {"text": "missing", "audio": "missing", "vision": "missing"},
        "modality_failure_reasons": {"text": reason, "audio": reason, "vision": reason},
        "vision_status": "missing",
        "vision_failure_reason": reason,
        "failure_reason": reason,
    }


def process_sample(
    config: Q1Config,
    entry: dict[str, Any],
    text_extractor: TextExtractor,
    asr_aligner: ASRAligner,
    audio_extractor: AudioExtractor,
    vision_extractor: VisionExtractor,
) -> dict[str, Any]:
    duration = float(entry["duration_alignment"])
    grid = make_time_grid(duration, config.aligned_length)
    sample_rate = int(config.get("audio_sample_rate", 16000))
    audio = extract_audio(Path(entry["video_path"]), sample_rate)
    text_raw = text_extractor.extract(entry["raw_text"])

    def align_candidate(candidate_words: list[dict[str, Any]]) -> list[dict[str, Any]]:
        return align_text_words(
            entry["raw_text"],
            candidate_words,
            duration,
            max_edit_ratio=float(config.get("text_match_max_edit_ratio", 0.40)),
            fallback_confidence=float(config.get("text_fallback_confidence", 0.20)),
            fallback_min_interval_s=float(config.get("text_fallback_min_interval_s", 0.01)),
        )

    def candidate_summary(candidate_words: list[dict[str, Any]], candidate_records: list[dict[str, Any]]) -> dict[str, Any]:
        fallback_count = int(sum(bool(item.get("fallback")) for item in candidate_records))
        invalid_count = int(sum(not bool(item.get("time_valid")) for item in candidate_records))
        exact_count = int(sum(item.get("match_type") == "exact" for item in candidate_records))
        edit_count = int(
            sum(item.get("match_type") in {"edit_distance", "equivalent_compound"} for item in candidate_records)
        )
        lexical_count = int(sum(bool(re.search(r"[A-Za-z0-9]", str(item.get("text", "")))) for item in candidate_words))
        return {
            "asr_word_count": int(len(candidate_words)),
            "lexical_asr_word_count": lexical_count,
            "fallback_count": fallback_count,
            "fallback_ratio": fallback_count / max(len(candidate_records), 1),
            "invalid_timestamp_count": invalid_count,
            "exact_match_count": exact_count,
            "edit_match_count": edit_count,
            "matched_count": int(len(candidate_records) - fallback_count),
        }

    baseline_profile = "baseline"
    asr_words = asr_aligner.extract(
        audio,
        duration,
        num_beams=int(config.get("asr_num_beams", 1)),
        temperature=float(config.get("asr_temperature", 0.0)),
        sample_rate=sample_rate,
    )
    aligned_words = align_candidate(asr_words)
    baseline_summary = candidate_summary(asr_words, aligned_words)
    selected_profile = baseline_profile
    retry_used = False
    retry_selected = False
    asr_attempts = [{"profile": baseline_profile, **baseline_summary}]
    retry_threshold = float(config.get("asr_retry_min_fallback_ratio", 0.50))
    retry_triggered = bool(
        config.get("asr_retry_enabled", True)
        and (
            baseline_summary["fallback_ratio"] >= retry_threshold
            or baseline_summary["lexical_asr_word_count"] == 0
        )
    )
    if retry_triggered:
        retry_profile = "retry_beam3_temperature0.2_context0.5"
        retry_words = asr_aligner.extract(
            audio,
            duration,
            num_beams=int(config.get("asr_retry_num_beams", 3)),
            temperature=float(config.get("asr_retry_temperature", 0.20)),
            context_s=float(config.get("asr_retry_context_s", 0.50)),
            sample_rate=sample_rate,
        )
        retry_aligned_words = align_candidate(retry_words)
        retry_summary = candidate_summary(retry_words, retry_aligned_words)
        asr_attempts.append({"profile": retry_profile, **retry_summary})
        retry_used = True
        if (
            retry_summary["fallback_count"] < baseline_summary["fallback_count"]
            and retry_summary["invalid_timestamp_count"] <= baseline_summary["invalid_timestamp_count"]
            and retry_summary["matched_count"] >= baseline_summary["matched_count"]
        ):
            asr_words = retry_words
            aligned_words = retry_aligned_words
            selected_profile = retry_profile
            retry_selected = True
    selected_summary = candidate_summary(asr_words, aligned_words)
    valid_text_positions = [
        position
        for position, index in enumerate(text_raw.word_indices)
        if aligned_words[index].get("time_valid")
        and aligned_words[index].get("start") is not None
        and aligned_words[index].get("end") is not None
    ]
    valid_text_indices = [text_raw.word_indices[position] for position in valid_text_positions]
    text_features = text_raw.features[valid_text_positions] if valid_text_positions else np.zeros((0, 768), dtype=np.float32)
    text_spans = np.asarray(
        [[aligned_words[index]["start"], aligned_words[index]["end"]] for index in valid_text_indices],
        dtype=np.float64,
    ).reshape((-1, 2))
    text_quality = np.asarray([aligned_words[index]["confidence"] for index in valid_text_indices], dtype=np.float64)
    text_aligned, text_mask, text_weights = overlap_pool(text_features, text_spans, grid, quality=text_quality)
    if len(text_features):
        text_confidence_aligned, _, _ = overlap_pool(
            text_quality[:, None].astype(np.float32), text_spans, grid, valid_mask=np.ones(len(text_features), dtype=bool)
        )
        text_confidence_aligned = text_confidence_aligned[:, 0]
    else:
        text_confidence_aligned = np.zeros(config.aligned_length, dtype=np.float32)
    text_confidence_aligned = np.clip(text_confidence_aligned, 0.0, 1.0).astype(np.float32)
    text_low_confidence_mask = (
        (text_mask.astype(bool))
        & (text_confidence_aligned < float(config.get("text_low_confidence_threshold", 0.50)))
    ).astype(np.uint8)

    audio_raw = audio_extractor.extract(audio, duration)
    audio_aligned, audio_mask, audio_weights = overlap_pool(
        audio_raw.features,
        audio_raw.spans,
        grid,
        quality=audio_raw.mask,
        valid_mask=audio_raw.mask.astype(bool),
    )
    prosody_aligned, prosody_mask, _ = overlap_pool(
        audio_raw.prosody,
        audio_raw.spans,
        grid,
        quality=audio_raw.prosody_mask,
        valid_mask=audio_raw.prosody_mask.astype(bool),
    )

    vision_raw = vision_extractor.extract(Path(entry["video_path"]), duration)
    vision_aligned, vision_mask, vision_weights = overlap_pool(
        vision_raw.features,
        vision_raw.spans,
        grid,
        quality=vision_raw.quality,
        valid_mask=vision_raw.mask.astype(bool),
    )
    emotion_aligned, emotion_mask, _ = overlap_pool(
        vision_raw.emotion_probs,
        vision_raw.spans,
        grid,
        quality=vision_raw.quality,
        valid_mask=vision_raw.mask.astype(bool),
    )
    vision_mask = np.minimum(vision_mask, emotion_mask).astype(np.uint8)

    text_fallback_count = int(selected_summary["fallback_count"])
    text_invalid_timestamp_count = int(selected_summary["invalid_timestamp_count"])
    text_fallback_ratio = text_fallback_count / max(len(aligned_words), 1)
    text_low_confidence_ratio = float(text_low_confidence_mask.mean())
    vision_valid_ratio = float(vision_mask.mean())
    vision_face_detection_ratio = float(vision_raw.detector_stats.get("face_detection_ratio", 0.0))
    vision_threshold = float(config.get("vision_degraded_threshold", 0.50))
    if vision_valid_ratio <= 0.0:
        vision_status, vision_failure_reason = "missing", "no_face_detected"
    elif vision_valid_ratio < vision_threshold:
        vision_status, vision_failure_reason = "degraded", "low_face_coverage"
    else:
        vision_status, vision_failure_reason = "ok", ""
    text_status = "missing" if not text_mask.any() else ("degraded" if text_low_confidence_mask.any() else "ok")
    audio_status = "missing" if not audio_mask.any() else "ok"
    modality_status = {"text": text_status, "audio": audio_status, "vision": vision_status}
    status_severity = {"ok": 0, "degraded": 1, "missing": 2}
    quality_status = max(modality_status.values(), key=lambda item: status_severity[item])
    source_counts: dict[str, int] = {}
    for item in aligned_words:
        source = str(item.get("timestamp_source", "unknown"))
        source_counts[source] = source_counts.get(source, 0) + 1
    text_timestamp_source = ";".join(f"{key}:{source_counts[key]}" for key in sorted(source_counts))
    lexical_asr_count = int(selected_summary["lexical_asr_word_count"])
    if text_fallback_ratio >= 0.999999:
        if lexical_asr_count == 0:
            text_review_status = "asr_empty_or_nonlexical_manual_review"
        else:
            text_review_status = "asr_text_mismatch_manual_review"
    elif text_fallback_ratio >= 0.50:
        text_review_status = "high_fallback_manual_review"
    elif text_invalid_timestamp_count:
        text_review_status = "partial_invalid_timestamp_manual_review"
    else:
        text_review_status = "normal"

    valid_lengths = {
        "text": int(text_mask.sum()),
        "audio": int(audio_mask.sum()),
        "vision": int(vision_mask.sum()),
    }
    stats = {
        "text_valid_ratio": float(text_mask.mean()),
        "text_low_confidence_ratio": text_low_confidence_ratio,
        "audio_valid_ratio": float(audio_mask.mean()),
        "vision_valid_ratio": vision_valid_ratio,
        "text_word_count": len(aligned_words),
        "text_asr_word_count": len(asr_words),
        "text_exact_match_count": int(selected_summary["exact_match_count"]),
        "text_edit_match_count": int(selected_summary["edit_match_count"]),
        "text_fallback_count": text_fallback_count,
        "text_fallback_ratio": text_fallback_ratio,
        "text_invalid_timestamp_count": text_invalid_timestamp_count,
        "text_timestamp_source": text_timestamp_source,
        "text_lexical_asr_word_count": lexical_asr_count,
        "text_review_status": text_review_status,
        "text_asr_retry_used": int(retry_used),
        "text_asr_retry_selected": int(retry_selected),
        "text_asr_selected_profile": selected_profile,
        "text_asr_attempt_count": len(asr_attempts),
        "audio_raw_frame_count": int(len(audio_raw.features)),
        "audio_decoded_duration": float(len(audio) / max(sample_rate, 1)),
        "vision_raw_frame_count": int(vision_raw.frame_count),
        "vision_decoded_frame_count": int(vision_raw.detector_stats.get("decoded_frames", vision_raw.frame_count)),
        "vision_discarded_out_of_timeline_frames": int(vision_raw.detector_stats.get("discarded_out_of_timeline_frames", 0)),
        "vision_face_frame_count": int(vision_raw.detector_stats["face_frames"]),
        "vision_face_detection_ratio": vision_face_detection_ratio,
        "vision_imputed_ratio": 0.0,
        "vision_track_switches": int(vision_raw.detector_stats["track_switches"]),
        "text_weight_sum": float(text_weights.sum()),
        "audio_weight_sum": float(audio_weights.sum()),
        "vision_weight_sum": float(vision_weights.sum()),
        "wavlm_effective_stride_samples": int(audio_raw.timing.get("effective_stride_samples", 0)),
        "wavlm_receptive_field_samples": int(audio_raw.timing.get("receptive_field_samples", 0)),
    }
    raw_audit = {
        "id": entry["id"],
        "text": {
            "features": text_features.astype(np.float16),
            "spans": text_spans.astype(np.float32),
            "words": aligned_words,
            "word_indices": valid_text_indices,
            "invalid_word_indices": [index for index, item in enumerate(aligned_words) if not item.get("time_valid")],
            "asr_words": asr_words,
            "asr_attempts": asr_attempts,
            "overflow_chunks": text_raw.overflow_chunks,
        },
        "audio": {
            "features": audio_raw.features.astype(np.float16),
            "spans": audio_raw.spans.astype(np.float32),
            "mask": audio_raw.mask,
            "prosody": audio_raw.prosody.astype(np.float32),
            "prosody_mask": audio_raw.prosody_mask,
            "timing": audio_raw.timing,
        },
        "vision": {
            "features": vision_raw.features.astype(np.float16),
            "emotion_probs": vision_raw.emotion_probs.astype(np.float32),
            "spans": vision_raw.spans.astype(np.float32),
            "mask": vision_raw.mask,
            "quality": vision_raw.quality,
            "face_bbox": vision_raw.face_bbox,
            "face_confidence": vision_raw.face_confidence,
            "track_ids": vision_raw.track_ids,
            "frame_indices": vision_raw.frame_indices,
        },
    }
    return {
        "text": text_aligned,
        "audio": audio_aligned,
        "vision": vision_aligned,
        "text_mask": text_mask,
        "text_confidence": text_confidence_aligned,
        "text_low_confidence_mask": text_low_confidence_mask,
        "audio_mask": audio_mask,
        "vision_mask": vision_mask,
        "vision_emotion_probs": emotion_aligned,
        "audio_prosody": prosody_aligned,
        "audio_prosody_mask": prosody_mask,
        "duration": duration,
        "duration_container": float(entry["duration_container"]),
        "duration_alignment": duration,
        "time_grid": grid,
        "raw_lengths": {"text": int(len(text_features)), "audio": int(len(audio_raw.features)), "vision": int(vision_raw.frame_count)},
        "valid_lengths": valid_lengths,
        "stats": stats,
        "raw_audit": raw_audit,
        "status": "complete",
        "processing_status": "complete",
        "quality_status": quality_status,
        "modality_status": modality_status,
        "modality_failure_reasons": {
            "text": "low_confidence_or_asr_mismatch" if text_status == "degraded" else "",
            "audio": "no_valid_audio_frames" if audio_status == "missing" else "",
            "vision": vision_failure_reason,
        },
        "vision_status": vision_status,
        "vision_failure_reason": vision_failure_reason,
        "failure_reason": "",
    }


def _write_audit_files(audit_dir: Path, result: dict[str, Any]) -> None:
    sample_id = result["raw_audit"]["id"]
    safe = _safe_name(sample_id)
    timestamp_dir = audit_dir / "raw_timestamps"
    face_dir = audit_dir / "face_tracks"
    timestamp_dir.mkdir(parents=True, exist_ok=True)
    face_dir.mkdir(parents=True, exist_ok=True)
    raw = result["raw_audit"]
    write_json(
        {
            "id": sample_id,
            "text_words": raw.get("text", {}).get("words", []),
            "asr_words": raw.get("text", {}).get("asr_words", []),
            "audio_spans": raw.get("audio", {}).get("spans", []),
            "vision_spans": raw.get("vision", {}).get("spans", []),
        },
        timestamp_dir / f"{safe}.json",
    )
    write_json(
        {
            "id": sample_id,
            "face_bbox": raw.get("vision", {}).get("face_bbox", []),
            "face_confidence": raw.get("vision", {}).get("face_confidence", []),
            "mask": raw.get("vision", {}).get("mask", []),
            "track_ids": raw.get("vision", {}).get("track_ids", []),
        },
        face_dir / f"{safe}.json",
    )


def _submission_readme(config: Q1Config, runtime: dict[str, Any], run_id: str) -> str:
    return f"""# Question 1 feature artifact

This directory is generated from the 100 raw videos in attachment 1. It is an independent Question 1 interface with three 768-dimensional main features:

```text
text   [N, 50, 768]
audio  [N, 50, 768]
vision [N, 50, 768]
```

The Question 2 and Question 3 models must continue to use attachment 2's `768/74/35` interface. This artifact is not a replacement for attachment 2.

The provided transcript is retained as `raw_text`. Whisper word timestamps are used only to locate the provided words in time. Text matching uses conservative normalization plus monotonic dynamic programming; high-fallback samples may receive one logged beam/temperature retry, and the retry is selected only when it improves coverage without increasing invalid timestamps. The 50 bins are proportional half-open intervals over `duration_alignment`, which follows the edit-list-aware FFmpeg presentation timeline. The source container/movie-header (`mvhd`) duration is retained separately as `duration_container` for auditing. Each feature is pooled by temporal overlap and quality weight. Invalid bins are zero after standardization and have mask value 0.

Run ID: `{run_id}`  
Config hash: `{config.config_hash}`  
Model profile: `{config.get('model_profile', 'uniform_768_reproducible')}`

Files:

- `q1_aligned_50.pkl`: submission candidate with aligned features, masks, lengths, timestamps and metadata.
- `manifest.csv`: one row per source sample.
- `alignment_log.csv`: extraction, alignment and quality statistics.
- `duration_audit.csv`: container, stream, decoded and selected alignment durations.
- `text_review.csv`: per-sample ASR lexical coverage, fallback ratio and manual-review status.
- `text_audit_summary.csv`: lightweight per-sample text conclusion for review and reuse.
- `normalization.json`: per-modality z-score parameters computed from valid Question 1 bins.
- `config.yaml`: resolved configuration and runtime versions.

`processing_status` describes whether extraction completed. `quality_status` and `vision_status` separately identify degraded or missing modalities. Text fallback words remain traceable in the audit files and are represented by `text_confidence` and `text_low_confidence_mask`.

`duration_alignment` is the video edit-list presentation axis. `duration_container` is retained for source-range checks. Positive decoded endpoint gaps are expected only within the recorded endpoint rule; the uncovered tail is not imputed and remains masked.

The compact reproducibility archive additionally contains the Question 1 source code, the source configuration, environment snapshots, execution entry points, and typical-sample figures/sidecars. The full `q1_unaligned.pkl`, model weights, source videos and caches remain outside the archive.
"""


PACKAGE_REQUIRED_PATHS = [
    "README.md",
    "package_manifest.json",
    "q1_submission/q1_aligned_50.pkl",
    "q1_submission/manifest.csv",
    "q1_submission/alignment_log.csv",
    "q1_submission/duration_audit.csv",
    "q1_submission/text_review.csv",
    "q1_submission/text_audit_summary.csv",
    "q1_submission/validation_report.json",
    "code/q1/pipeline.py",
    "code/q1/validate.py",
    "code/q1/utils.py",
    "config/q1.yaml",
    "environment/environment.json",
    "environment/pip_freeze.txt",
    "tools/run_q1.sh",
    "tools/validate_q1.sh",
    "audit/typical_samples.json",
]


def _prepare_submission_package(
    run_dir: Path,
    submission_dir: Path,
    audit_dir: Path,
    environment_dir: Path,
    config: Q1Config,
    run_id: str,
) -> Path:
    package_root = run_dir / "outputs" / ".q1_package_staging"
    if package_root.exists():
        shutil.rmtree(package_root)
    package_root.mkdir(parents=True, exist_ok=True)

    for source in submission_dir.rglob("*"):
        if source.is_file():
            target = package_root / "q1_submission" / source.relative_to(submission_dir)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    for source in config.project_root.joinpath("q1").rglob("*.py"):
        if "__pycache__" in source.parts:
            continue
        target = package_root / "code" / "q1" / source.relative_to(config.project_root / "q1")
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)

    config_target = package_root / "config" / config.path.name
    config_target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(config.path, config_target)

    for source in environment_dir.glob("*"):
        if source.is_file():
            target = package_root / "environment" / source.name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    for name in ("run_q1.sh", "run_q1.py", "validate_q1.sh"):
        source = config.project_root / "tools" / name
        if source.exists():
            target = package_root / "tools" / name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    for source in (audit_dir.iterdir() if audit_dir.exists() else []):
        if not source.is_file():
            continue
        if source.suffix.lower() in {".png", ".json"} and (
            source.name.startswith("typical_sample") or source.name == "typical_samples.json"
        ):
            target = package_root / "audit" / source.name
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    (package_root / "README.md").write_text(
        f"""# Question 1 reproducibility package

Run ID: `{run_id}`

This archive contains the Question 1 aligned feature output, the source code and configuration used to generate it, the remote environment snapshot, execution entry points, and typical-sample audit figures with sidecars.

The alignment axis is the video edit-list presentation timeline. The container duration, decoded endpoints, endpoint differences, masks, fallback records and quality statuses are retained in `q1_submission/`.

The full unaligned audit pickle, model weights, source videos and caches are intentionally excluded to keep the archive compact and independently reviewable.
""",
        encoding="utf-8",
    )
    write_json(
        {
            "schema_version": "q1.reproducibility_package.v1",
            "run_id": run_id,
            "required_paths": PACKAGE_REQUIRED_PATHS,
            "excluded": ["q1_unaligned.pkl", "models", "source videos", "caches"],
        },
        package_root / "package_manifest.json",
    )
    return package_root


def _zip_submission(package_root: Path, zip_path: Path) -> int:
    if zip_path.exists():
        zip_path.unlink()
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in sorted(package_root.rglob("*")):
            if path.is_file():
                archive.write(path, path.relative_to(package_root))
    return zip_path.stat().st_size


def _build_submission_zip(
    run_dir: Path,
    submission_dir: Path,
    audit_dir: Path,
    environment_dir: Path,
    config: Q1Config,
    run_id: str,
    zip_path: Path,
) -> int:
    package_root = _prepare_submission_package(
        run_dir,
        submission_dir,
        audit_dir,
        environment_dir,
        config,
        run_id,
    )
    try:
        return _zip_submission(package_root, zip_path)
    finally:
        shutil.rmtree(package_root, ignore_errors=True)


def run_pipeline(
    config: Q1Config,
    run_id: str,
    limit: int | None = None,
    sample_id: str | None = None,
    allow_failures: bool = False,
    skip_model_hash: bool = False,
) -> Path:
    run_dir = config.path_for("run_root") / run_id
    if run_dir.exists():
        raise FileExistsError(f"Run directory already exists: {run_dir}")
    submission_dir = run_dir / "outputs" / "q1_submission"
    audit_dir = run_dir / "outputs" / "q1_audit"
    log_path = run_dir / "logs" / f"q1_{datetime.now().strftime('%Y%m%d_%H%M%S')}.log"
    run_dir.mkdir(parents=True, exist_ok=False)
    logger = configure_logging(log_path)
    (run_dir / "RUNNING").write_text("", encoding="utf-8")
    started = time.time()
    try:
        seed = int(config.get("random_seed", 2026))
        set_deterministic(seed)
        logger.info("Starting Question 1 run %s", run_id)
        environment = environment_snapshot(config.path_for("model_manifest"))
        write_json(environment, run_dir / "environment" / "environment.json")
        pip_freeze = subprocess.run(
            [sys.executable, "-m", "pip", "freeze"],
            check=False,
            capture_output=True,
            text=True,
        )
        (run_dir / "environment" / "pip_freeze.txt").write_text(
            pip_freeze.stdout + ("\n" + pip_freeze.stderr if pip_freeze.stderr else ""),
            encoding="utf-8",
        )
        model_check = verify_model_manifest(config, logger, verify_hashes=not skip_model_hash)
        write_json(model_check, run_dir / "environment" / "model_verification.json")
        manifest = build_manifest(config, logger)
        if sample_id:
            manifest = [entry for entry in manifest if entry["id"] == sample_id]
            if not manifest:
                raise KeyError(f"Unknown sample ID: {sample_id}")
        if limit is not None:
            manifest = manifest[: int(limit)]
        logger.info("Samples selected for run: %d", len(manifest))

        import torch

        device = torch.device("cuda:0" if torch.cuda.is_available() else "cpu")
        logger.info("Inference device: %s", device)
        models_root = config.path_for("models_root")
        text_extractor = TextExtractor(config.path_for("text_model_path"), device, int(config.get("text_max_length", 8192)))
        asr_aligner = ASRAligner(
            config.path_for("asr_timestamp_model_path"),
            device,
            float(config.get("asr_window_s", 30.0)),
            float(config.get("asr_overlap_s", 5.0)),
        )
        audio_extractor = AudioExtractor(
            config.path_for("audio_model_path"), device, int(config.get("audio_sample_rate", 16000))
        )
        vision_extractor = VisionExtractor(
            config.path_for("vision_model_path"),
            device,
            float(config.get("video_output_fps", 25.0)),
            float(config.get("face_detector_min_confidence", 0.5)),
            float(config.get("face_crop_margin", 0.20)),
            float(config.get("face_track_iou_threshold", 0.20)),
            int(config.get("face_track_max_gap_frames", 2)),
        )
        del models_root

        results: list[dict[str, Any]] = []
        audit_results: list[dict[str, Any]] = []
        completed = 0
        failures = 0
        for index, entry in enumerate(manifest, start=1):
            logger.info("[%d/%d] Processing %s", index, len(manifest), entry["id"])
            try:
                result = process_sample(config, entry, text_extractor, asr_aligner, audio_extractor, vision_extractor)
                completed += 1
                logger.info(
                    "Completed %s: text=%.3f audio=%.3f vision=%.3f",
                    entry["id"],
                    result["stats"]["text_valid_ratio"],
                    result["stats"]["audio_valid_ratio"],
                    result["stats"]["vision_valid_ratio"],
                )
            except Exception as exc:  # preserve a traceable row for every source sample
                failures += 1
                reason = f"{type(exc).__name__}: {exc}"
                logger.error("Failed %s: %s\n%s", entry["id"], reason, traceback.format_exc())
                result = _empty_result(config, entry, reason)
            _write_audit_files(audit_dir, result)
            result["raw_audit"]["entry"] = {key: value for key, value in entry.items() if key != "video_path"}
            results.append(result)
            audit_results.append(result["raw_audit"])

        if not results:
            raise RuntimeError("No samples were processed")
        text_stack = np.stack([item["text"] for item in results]).astype(np.float32)
        audio_stack = np.stack([item["audio"] for item in results]).astype(np.float32)
        vision_stack = np.stack([item["vision"] for item in results]).astype(np.float32)
        text_masks = np.stack([item["text_mask"] for item in results])
        audio_masks = np.stack([item["audio_mask"] for item in results])
        vision_masks = np.stack([item["vision_mask"] for item in results])
        text_stack, text_mean, text_std, _ = standardize_aligned(text_stack, text_masks)
        audio_stack, audio_mean, audio_std, _ = standardize_aligned(audio_stack, audio_masks)
        vision_stack, vision_mean, vision_std, _ = standardize_aligned(vision_stack, vision_masks)
        submission_dir.mkdir(parents=True, exist_ok=True)
        artifact = {
            "schema_version": "q1.uniform_768.v1",
            "model_profile": config.get("model_profile", "uniform_768_reproducible"),
            "config_hash": config.config_hash,
            "id": [entry["id"] for entry in manifest],
            "text": text_stack.astype(np.float16),
            "audio": audio_stack.astype(np.float16),
            "vision": vision_stack.astype(np.float16),
            "auxiliary": {
                "vision_emotion_probs": np.stack([item["vision_emotion_probs"] for item in results]).astype(np.float32),
                "audio_prosody": np.stack([item["audio_prosody"] for item in results]).astype(np.float32),
                "audio_prosody_mask": np.stack([item["audio_prosody_mask"] for item in results]).astype(np.uint8),
            },
            "masks": {
                "text": text_masks.astype(np.uint8),
                "audio": audio_masks.astype(np.uint8),
                "vision": vision_masks.astype(np.uint8),
            },
            "text_confidence": np.stack([item["text_confidence"] for item in results]).astype(np.float32),
            "text_low_confidence_mask": np.stack([item["text_low_confidence_mask"] for item in results]).astype(np.uint8),
            "durations": np.asarray([item["duration"] for item in results], dtype=np.float32),
            "duration_alignment": np.asarray([item["duration_alignment"] for item in results], dtype=np.float32),
            "duration_container": np.asarray([item["duration_container"] for item in results], dtype=np.float32),
            "time_grid": np.stack([item["time_grid"] for item in results]).astype(np.float32),
            "valid_lengths": {
                modality: np.asarray([item["valid_lengths"][modality] for item in results], dtype=np.int16)
                for modality in ("text", "audio", "vision")
            },
            "raw_lengths": {
                modality: np.asarray([item["raw_lengths"][modality] for item in results], dtype=np.int32)
                for modality in ("text", "audio", "vision")
            },
            "raw_text": [entry["raw_text"] for entry in manifest],
            "labels": np.asarray([entry["label"] if entry["label"] is not None else np.nan for entry in manifest], dtype=np.float32),
            "annotations": [entry["annotation"] for entry in manifest],
            "status": [item["status"] for item in results],
            "processing_status": [item["processing_status"] for item in results],
            "quality_status": [item["quality_status"] for item in results],
            "modality_status": [item["modality_status"] for item in results],
            "modality_failure_reasons": [item["modality_failure_reasons"] for item in results],
            "vision_status": [item["vision_status"] for item in results],
            "vision_failure_reasons": [item["vision_failure_reason"] for item in results],
            "failure_reasons": [item["failure_reason"] for item in results],
            "metadata": [
                {
                    "video_id": entry["video_id"],
                    "clip_id": entry["clip_id"],
                    "video_path": entry["video_path_relative"],
                    "sha256": entry["sha256"],
                    "duration_container": entry["duration_container"],
                    "duration_alignment": entry["duration_alignment"],
                    "duration_source": entry["duration_source"],
                    "text_fallback_ratio": result["stats"]["text_fallback_ratio"],
                    "text_review_status": result["stats"]["text_review_status"],
                    "vision_valid_ratio": result["stats"]["vision_valid_ratio"],
                    "vision_status": result["vision_status"],
                    "quality_status": result["quality_status"],
                    "processing_status": result["processing_status"],
                }
                for entry, result in zip(manifest, results)
            ],
        }
        artifact_path = submission_dir / "q1_aligned_50.pkl"
        atomic_pickle_dump(artifact, artifact_path)
        normalization = {
            "method": "zscore_on_valid_aligned_bins",
            "epsilon": 1e-6,
            "text": {"mean": text_mean, "std": text_std, "dim": 768},
            "audio": {"mean": audio_mean, "std": audio_std, "dim": 768},
            "vision": {"mean": vision_mean, "std": vision_std, "dim": 768},
        }
        write_json(normalization, submission_dir / "normalization.json")

        manifest_rows = []
        alignment_rows = []
        for entry, result in zip(manifest, results):
            entry = dict(entry)
            entry.update(
                {
                    "text_length": result["raw_lengths"]["text"],
                    "audio_length": result["raw_lengths"]["audio"],
                    "vision_length": result["raw_lengths"]["vision"],
                    "feature_dim_text": 768,
                    "feature_dim_audio": 768,
                    "feature_dim_vision": 768,
                    "aligned_length": config.aligned_length,
                    "alignment_granularity": "proportional_50_half_open_intervals",
                    "feature_status": result["processing_status"],
                    "processing_status": result["processing_status"],
                    "quality_status": result["quality_status"],
                    "text_status": result["modality_status"]["text"],
                    "audio_status": result["modality_status"]["audio"],
                    "vision_status": result["vision_status"],
                    "vision_failure_reason": result["vision_failure_reason"],
                    "vision_valid_ratio": result["stats"]["vision_valid_ratio"],
                    "vision_face_detection_ratio": result["stats"]["vision_face_detection_ratio"],
                    "vision_imputed_ratio": result["stats"]["vision_imputed_ratio"],
                    "text_fallback_count": result["stats"]["text_fallback_count"],
                    "text_fallback_ratio": result["stats"]["text_fallback_ratio"],
                    "text_word_count": result["stats"]["text_word_count"],
                    "text_asr_word_count": result["stats"]["text_asr_word_count"],
                    "text_invalid_timestamp_count": result["stats"]["text_invalid_timestamp_count"],
                    "text_low_confidence_ratio": result["stats"]["text_low_confidence_ratio"],
                    "text_timestamp_source": result["stats"]["text_timestamp_source"],
                    "text_lexical_asr_word_count": result["stats"]["text_lexical_asr_word_count"],
                    "text_review_status": result["stats"]["text_review_status"],
                    "text_asr_retry_used": result["stats"]["text_asr_retry_used"],
                    "text_asr_retry_selected": result["stats"]["text_asr_retry_selected"],
                    "text_asr_selected_profile": result["stats"]["text_asr_selected_profile"],
                    "text_asr_attempt_count": result["stats"]["text_asr_attempt_count"],
                    "failure_reason": result["failure_reason"],
                }
            )
            manifest_rows.append(entry)
            alignment_row = {
                "id": entry["id"],
                "duration_container": entry["duration_container"],
                "duration_alignment": entry["duration_alignment"],
                "duration_source": entry["duration_source"],
                "duration_alignment_minus_decoded_video": entry["duration_alignment_minus_decoded_video"],
                "duration_alignment_minus_decoded_audio": entry["duration_alignment_minus_decoded_audio"],
                "duration_alignment_minus_decoded_max": entry["duration_alignment_minus_decoded_max"],
                "duration_endpoint_status": entry["duration_endpoint_status"],
                "duration_endpoint_issues": entry["duration_endpoint_issues"],
                "processing_status": result["processing_status"],
                "quality_status": result["quality_status"],
                "failure_reason": result["failure_reason"],
            }
            alignment_row.update(result["stats"])
            alignment_rows.append(alignment_row)
        manifest_fields = [
            "id", "video_id", "clip_id", "video_path_relative", "raw_text", "label", "annotation",
            "duration_container", "duration_format", "duration_mvhd", "duration_video_stream", "duration_audio_stream",
            "duration_video_edit", "duration_audio_edit", "duration_decoded_video", "duration_decoded_audio",
            "duration_decoded", "duration_alignment", "duration_source", "duration_start_time",
            "video_edit_list_entries", "audio_edit_list_entries", "video_edit_list_present", "audio_edit_list_present",
            "duration_decoded_video_start_pts", "duration_decoded_video_last_pts", "duration_decoded_video_end_pts",
            "duration_decoded_video_packet_count", "duration_decoded_video_packet_duration",
            "duration_decoded_audio_start_pts", "duration_decoded_audio_last_pts", "duration_decoded_audio_end_pts",
            "duration_decoded_audio_packet_count", "duration_decoded_audio_packet_duration",
            "duration_alignment_minus_decoded_video", "duration_alignment_minus_decoded_audio",
            "duration_alignment_minus_decoded_max", "duration_endpoint_video_tolerance",
            "duration_endpoint_audio_tolerance", "duration_endpoint_negative_tolerance",
            "duration_endpoint_status", "duration_endpoint_issues", "duration_endpoint_rule",
            "duration_endpoint_explanation",
            "audio_sample_rate_original", "video_fps", "video_width", "video_height", "text_length", "audio_length",
            "vision_length", "feature_dim_text", "feature_dim_audio", "feature_dim_vision", "aligned_length",
            "alignment_granularity", "feature_status", "processing_status", "quality_status", "text_status",
            "audio_status", "vision_status", "vision_failure_reason", "vision_valid_ratio",
            "vision_face_detection_ratio", "vision_imputed_ratio", "text_fallback_count", "text_fallback_ratio",
            "text_word_count", "text_asr_word_count", "text_invalid_timestamp_count",
            "text_low_confidence_ratio", "text_timestamp_source", "failure_reason", "sha256",
            "text_lexical_asr_word_count", "text_review_status", "text_asr_retry_used",
            "text_asr_retry_selected", "text_asr_selected_profile", "text_asr_attempt_count",
        ]
        write_csv(manifest_rows, submission_dir / "manifest.csv", manifest_fields)
        alignment_fields = sorted({key for row in alignment_rows for key in row})
        write_csv(alignment_rows, submission_dir / "alignment_log.csv", alignment_fields)
        write_csv(
            [
                {
                    "id": entry["id"],
                    "raw_text": entry["raw_text"],
                    "asr_word_count": result["stats"]["text_asr_word_count"],
                    "lexical_asr_word_count": result["stats"]["text_lexical_asr_word_count"],
                    "fallback_count": result["stats"]["text_fallback_count"],
                    "fallback_ratio": result["stats"]["text_fallback_ratio"],
                    "invalid_timestamp_count": result["stats"]["text_invalid_timestamp_count"],
                    "review_status": result["stats"]["text_review_status"],
                    "asr_retry_used": result["stats"]["text_asr_retry_used"],
                    "asr_retry_selected": result["stats"]["text_asr_retry_selected"],
                    "asr_selected_profile": result["stats"]["text_asr_selected_profile"],
                    "asr_attempt_count": result["stats"]["text_asr_attempt_count"],
                    "asr_words": result["raw_audit"].get("text", {}).get("asr_words", []),
                }
                for entry, result in zip(manifest, results)
            ],
            submission_dir / "text_review.csv",
            [
                "id", "raw_text", "asr_word_count", "lexical_asr_word_count", "fallback_count",
                "fallback_ratio", "invalid_timestamp_count", "review_status", "asr_retry_used",
                "asr_retry_selected", "asr_selected_profile", "asr_attempt_count", "asr_words",
            ],
        )
        write_csv(
            [
                {
                    "id": entry["id"],
                    "word_count": result["stats"]["text_word_count"],
                    "asr_word_count": result["stats"]["text_asr_word_count"],
                    "lexical_asr_word_count": result["stats"]["text_lexical_asr_word_count"],
                    "fallback_count": result["stats"]["text_fallback_count"],
                    "fallback_ratio": result["stats"]["text_fallback_ratio"],
                    "invalid_timestamp_count": result["stats"]["text_invalid_timestamp_count"],
                    "review_status": result["stats"]["text_review_status"],
                    "manual_review_required": int(result["stats"]["text_review_status"] != "normal"),
                    "asr_retry_used": result["stats"]["text_asr_retry_used"],
                    "asr_retry_selected": result["stats"]["text_asr_retry_selected"],
                    "asr_selected_profile": result["stats"]["text_asr_selected_profile"],
                    "asr_attempt_count": result["stats"]["text_asr_attempt_count"],
                    "conclusion": _text_review_conclusion(result["stats"]),
                }
                for entry, result in zip(manifest, results)
            ],
            submission_dir / "text_audit_summary.csv",
            [
                "id", "word_count", "asr_word_count", "lexical_asr_word_count", "fallback_count",
                "fallback_ratio", "invalid_timestamp_count", "review_status", "manual_review_required",
                "asr_retry_used", "asr_retry_selected", "asr_selected_profile", "asr_attempt_count", "conclusion",
            ],
        )
        duration_fields = [
            "id", "duration_container", "duration_format", "duration_mvhd", "duration_video_stream", "duration_audio_stream",
            "duration_video_edit", "duration_audio_edit", "duration_decoded_video", "duration_decoded_audio",
            "duration_decoded", "duration_alignment", "duration_source", "duration_start_time",
            "video_edit_list_entries", "audio_edit_list_entries", "video_edit_list_present", "audio_edit_list_present",
            "duration_decoded_video_start_pts", "duration_decoded_video_last_pts", "duration_decoded_video_end_pts",
            "duration_decoded_video_packet_count", "duration_decoded_video_packet_duration",
            "duration_decoded_audio_start_pts", "duration_decoded_audio_last_pts", "duration_decoded_audio_end_pts",
            "duration_decoded_audio_packet_count", "duration_decoded_audio_packet_duration",
            "duration_alignment_minus_decoded_video", "duration_alignment_minus_decoded_audio",
            "duration_alignment_minus_decoded_max", "duration_endpoint_video_tolerance",
            "duration_endpoint_audio_tolerance", "duration_endpoint_negative_tolerance",
            "duration_endpoint_status", "duration_endpoint_issues", "duration_endpoint_rule",
            "duration_endpoint_explanation",
        ]
        write_csv(manifest_rows, submission_dir / "duration_audit.csv", duration_fields)

        if bool(config.get("save_unaligned", True)):
            atomic_pickle_dump({"schema_version": "q1.unaligned.v1", "samples": audit_results}, audit_dir / "q1_unaligned.pkl")
        write_json(environment, audit_dir / "environment.json")
        resolved_config = dict(config.values)
        resolved_config["runtime"] = environment
        resolved_config["model_verification"] = model_check
        resolved_config["config_hash"] = config.config_hash
        dump_yaml(resolved_config, submission_dir / "config.yaml")
        (submission_dir / "README.md").write_text(_submission_readme(config, environment, run_id), encoding="utf-8")

        if bool(config.get("generate_typical_sample_figure", True)):
            try:
                from .visualize_sample import render_sample_figure

                quality = np.minimum.reduce(
                    [text_masks.mean(axis=1), audio_masks.mean(axis=1), vision_masks.mean(axis=1)]
                )
                id_to_index = {entry["id"]: index for index, entry in enumerate(manifest)}
                requested_ids = [str(item) for item in config.get("audit_sample_ids", [])]
                selected_indices = [id_to_index[item] for item in requested_ids if item in id_to_index]
                if not selected_indices:
                    selected_indices = [int(np.argmax(quality))]
                sample_summaries = []
                for audit_position, sample_index in enumerate(selected_indices):
                    sample_id = manifest[sample_index]["id"]
                    safe_sample_id = _safe_name(sample_id)
                    aligned_view = {
                        "time_grid": artifact["time_grid"][sample_index],
                        "text": artifact["text"][sample_index],
                        "audio": artifact["audio"][sample_index],
                        "vision": artifact["vision"][sample_index],
                        "masks": {key: value[sample_index] for key, value in artifact["masks"].items()},
                        "vision_emotion_probs": artifact["auxiliary"]["vision_emotion_probs"][sample_index],
                        "text_confidence": artifact["text_confidence"][sample_index],
                        "text_low_confidence_mask": artifact["text_low_confidence_mask"][sample_index],
                    }
                    figure_path = audit_dir / f"typical_sample_{safe_sample_id}_alignment.png"
                    render_sample_figure(
                        sample_id,
                        aligned_view,
                        figure_path,
                        raw_audit=audit_results[sample_index],
                        source_path=Path(manifest[sample_index]["video_path"]),
                        duration_container=float(manifest[sample_index]["duration_container"]),
                        duration_alignment=float(manifest[sample_index]["duration_alignment"]),
                    )
                    if audit_position == 0:
                        shutil.copy2(figure_path, audit_dir / "typical_sample_alignment.png")
                    raw_sample = audit_results[sample_index]
                    sidecar = {
                        "id": sample_id,
                        "index": sample_index,
                        "duration_container": manifest[sample_index]["duration_container"],
                        "duration_alignment": manifest[sample_index]["duration_alignment"],
                        "duration_source": manifest[sample_index]["duration_source"],
                        "time_grid": artifact["time_grid"][sample_index],
                        "text_words": raw_sample.get("text", {}).get("words", []),
                        "asr_words": raw_sample.get("text", {}).get("asr_words", []),
                        "audio_spans": raw_sample.get("audio", {}).get("spans", []),
                        "audio_mask": raw_sample.get("audio", {}).get("mask", []),
                        "vision_spans": raw_sample.get("vision", {}).get("spans", []),
                        "vision_mask": raw_sample.get("vision", {}).get("mask", []),
                        "vision_frame_indices": raw_sample.get("vision", {}).get("frame_indices", []),
                        "face_bbox": raw_sample.get("vision", {}).get("face_bbox", []),
                        "aligned_masks": {key: value[sample_index] for key, value in artifact["masks"].items()},
                        "text_confidence": artifact["text_confidence"][sample_index],
                        "text_low_confidence_mask": artifact["text_low_confidence_mask"][sample_index],
                    }
                    write_json(sidecar, audit_dir / f"typical_sample_{safe_sample_id}_sidecar.json")
                    sample_summaries.append(
                        {
                            "id": sample_id,
                            "index": sample_index,
                            "combined_valid_ratio": float(quality[sample_index]),
                            "figure": figure_path.name,
                            "sidecar": f"typical_sample_{safe_sample_id}_sidecar.json",
                            "quality_status": results[sample_index]["quality_status"],
                        }
                    )
                write_json(sample_summaries, audit_dir / "typical_samples.json")
                write_json(sample_summaries[0], audit_dir / "typical_sample.json")
            except Exception as exc:
                logger.warning("Typical sample visualization was not generated: %s", exc)

        from .validate import validate_artifact, validate_submission_package

        raw_audit_path = audit_dir / "q1_unaligned.pkl"
        validation = validate_artifact(
            artifact_path,
            expected_count=len(manifest),
            expected_dims=config.main_feature_dims,
            require_complete=bool(config.get("require_all_samples_complete", True)) and not allow_failures,
            manifest_path=submission_dir / "manifest.csv",
            raw_audit_path=raw_audit_path if raw_audit_path.exists() else None,
            alignment_log_path=submission_dir / "alignment_log.csv",
            text_review_path=submission_dir / "text_review.csv",
            duration_audit_path=submission_dir / "duration_audit.csv",
        )
        write_json(validation, submission_dir / "validation_report.json")
        if validation["errors"]:
            raise RuntimeError("Output validation failed: " + "; ".join(validation["errors"]))

        zip_path = run_dir / "outputs" / f"q1_submission_{run_id}.zip"
        zip_bytes = _build_submission_zip(
            run_dir,
            submission_dir,
            audit_dir,
            run_dir / "environment",
            config,
            run_id,
            zip_path,
        )
        package_report = validate_submission_package(zip_path, PACKAGE_REQUIRED_PATHS)
        validation["package"] = package_report
        validation["conclusions"]["reproducibility_package"] = "passed" if not package_report["errors"] else "failed"
        validation["conclusions"]["overall"] = (
            "passed_with_quality_conditions"
            if validation["conclusions"].get("structure") == "passed" and not package_report["errors"]
            else "blocked"
        )
        validation["errors"].extend(package_report["errors"])
        validation["warnings"].extend(package_report["warnings"])
        write_json(validation, submission_dir / "validation_report.json")
        if validation["errors"]:
            raise RuntimeError("Submission package validation failed: " + "; ".join(validation["errors"]))

        # Rebuild so the final archive contains the final validation report.
        zip_bytes = _build_submission_zip(
            run_dir,
            submission_dir,
            audit_dir,
            run_dir / "environment",
            config,
            run_id,
            zip_path,
        )
        final_package_report = validate_submission_package(zip_path, PACKAGE_REQUIRED_PATHS)
        if final_package_report["errors"]:
            raise RuntimeError("Final submission package validation failed: " + "; ".join(final_package_report["errors"]))
        write_json(
            {
                "uncompressed_submission_bytes": sum(path.stat().st_size for path in submission_dir.rglob("*") if path.is_file()),
                "compressed_bytes": zip_bytes,
                "limit_bytes": int(config.get("submission_max_bytes", 52_428_800)),
                "within_limit": zip_bytes <= int(config.get("submission_max_bytes", 52_428_800)),
                "package_file_count": final_package_report["file_count"],
                "required_paths": PACKAGE_REQUIRED_PATHS,
            },
            submission_dir / "package_size.json",
        )
        if zip_bytes > int(config.get("submission_max_bytes", 52_428_800)):
            raise RuntimeError(f"Submission archive exceeds limit: {zip_bytes} bytes")

        elapsed = time.time() - started
        logger.info("Run complete: successful=%d failed=%d elapsed=%.1fs", completed, failures, elapsed)
        write_json(
            {
                "run_id": run_id,
                "started_at": datetime.fromtimestamp(started, timezone.utc).isoformat(),
                "finished_at": datetime.now(timezone.utc).isoformat(),
                "successful_samples": completed,
                "failed_samples": failures,
                "elapsed_seconds": elapsed,
                "submission_zip": str(zip_path),
            },
            run_dir / "run_summary.json",
        )
        (run_dir / "RUNNING").unlink(missing_ok=True)
        (run_dir / "DONE").write_text("", encoding="utf-8")
        return run_dir
    except Exception:
        logger.error("Run failed:\n%s", traceback.format_exc())
        (run_dir / "RUNNING").unlink(missing_ok=True)
        (run_dir / "FAILED").write_text(traceback.format_exc(), encoding="utf-8")
        raise


def main() -> None:
    parser = argparse.ArgumentParser(description="Run Question 1 raw-video feature extraction remotely.")
    parser.add_argument("--config", default="configs/q1.yaml")
    parser.add_argument("--run-id", required=True)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--sample-id")
    parser.add_argument("--allow-failures", action="store_true")
    parser.add_argument("--skip-model-hash", action="store_true")
    args = parser.parse_args()
    config = load_config(args.config)
    run_pipeline(
        config,
        args.run_id,
        limit=args.limit,
        sample_id=args.sample_id,
        allow_failures=args.allow_failures,
        skip_model_hash=args.skip_model_hash,
    )


if __name__ == "__main__":
    main()
